#!/usr/bin/env python3
"""Build a teacher-facing review report from scan results and agent decisions."""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CATEGORY_LABELS = {
    "self_harm": "自伤/自杀风险",
    "emotion": "情绪困扰",
    "violence": "霸凌/暴力",
    "family": "家庭风险",
    "sexual": "性侵害/骚扰",
    "grief": "丧亲/死亡",
    "harm": "伤害/受伤",
}
SEVERITY_LABELS = {"high": "高", "medium": "中", "low": "低", "uncertain": "待确认"}


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def known(value: object) -> bool:
    return value is not None


def display(value: object, suffix: str = "") -> str:
    if value is None:
        return "待确认"
    return f"{value}{suffix}"


def set_font(run, east: str = "宋体", size: float = 12, bold: bool = False):
    run.font.name = east
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), east)


def add_text(doc, text, *, east="宋体", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
             before=0, after=6, line=1.5):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    run = p.add_run(text)
    set_font(run, east, size, bold)
    return p


def add_heading(doc, text):
    return add_text(doc, text, east="黑体", size=14, bold=True, before=10, after=6)


def table_cell(cell, text, *, bold=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text))
    set_font(run, "宋体", 10.5, bold)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, title in enumerate(headers):
        table_cell(table.rows[0].cells[i], title, bold=True, center=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            table_cell(cells[i], value, center=i in (0, 3, 4))
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    return table


def normalize_decisions(decisions: dict, scan: dict) -> tuple[list[dict], list[dict], list[str]]:
    warnings = []
    if not decisions.get("reviewed"):
        warnings.append("decisions.reviewed 不是 true，报告仍为草稿，敏感结论未完成复核")
    confirmed = decisions.get("confirmed") or []
    excluded = decisions.get("excluded") or []
    if not isinstance(confirmed, list) or not isinstance(excluded, list):
        raise ValueError("sensitive_decisions.json 的 confirmed/excluded 必须是数组")

    valid_names = {str(item.get("name", "")) for item in scan.get("students", [])}
    normalized = []
    for index, item in enumerate(confirmed, start=1):
        if not isinstance(item, dict) or not item.get("student"):
            warnings.append(f"confirmed[{index}] 缺少 student，已跳过")
            continue
        student = str(item["student"]).strip()
        if valid_names and student not in valid_names:
            warnings.append(f"confirmed[{index}] 学生 {student} 不在扫描学生列表中")
        category = str(item.get("category") or "other")
        try:
            occurrences = max(1, int(item.get("occurrences") or 1))
        except (TypeError, ValueError):
            occurrences = 1
        normalized.append(
            {
                "student": student,
                "student_id": str(item.get("student_id") or ""),
                "category": category,
                "category_label": CATEGORY_LABELS.get(category, category),
                "summary": str(item.get("summary") or "未填写摘要").strip(),
                "label": str(item.get("label") or "").strip(),
                "quote": str(item.get("quote") or "").strip(),
                "context_type": str(item.get("context_type") or "uncertain").strip(),
                "severity": str(item.get("severity") or "uncertain").strip().lower(),
                "source_file": str(item.get("source_file") or "").strip(),
                "occurrences": occurrences,
            }
        )
    return normalized, excluded, warnings


def build_short_message(class_name: str, teacher: str, total, warning_count, missing_count,
                        completed_count, confirmed: list[dict], counts_unknown: bool) -> str:
    lines = [f"{teacher}，我的审核已完成。", ""]
    if counts_unknown:
        lines.append(f"{class_name}自传共{display(total)}人；预警、缺页漏页及完成情况待确认。")
    else:
        lines.append(
            f"{class_name}自传共{total}人，其中预警对象{display(warning_count, '位')}，"
            f"缺页漏页对象共{display(missing_count, '位')}，"
            f"符合要求并校对完成的对象共{display(completed_count, '位')}"
            "（包含被预警的同学）。"
        )
    lines.append("")
    if confirmed:
        students = []
        for item in confirmed:
            if item["student"] not in students:
                students.append(item["student"])
        lines.append(f"敏感内容涉及{len(students)}名学生：")
        for student in students:
            items = [item for item in confirmed if item["student"] == student]
            details = []
            for item in items:
                suffix = f"（{item['occurrences']}处）" if item.get("occurrences", 1) > 1 else ""
                text = item.get("label") or item["summary"]
                details.append(f"{text}{suffix}")
            lines.append(f"- {student}：{'；'.join(details)}")
    else:
        lines.append("本次未确认需在汇报中单列的敏感内容。")
    lines.extend(["", "审核后的文档已按班级归档，供老师查阅。"])
    return "\n".join(lines)


def build_markdown(class_name, teacher, total, warning_count, missing_count, completed_count,
                   warning_names, missing_names, confirmed, excluded, counts_unknown, warnings,
                   scan: dict, decisions: dict) -> str:
    lines = [
        f"# {class_name}自传审核汇总报告",
        "",
        f"> 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}  ",
        f"> 审核对象：`{scan.get('input_dir', '')}`",
        "",
        "## 一、总体情况",
        "",
        "| 项目 | 数量 |",
        "| --- | --- |",
        f"| 审核文件 | {scan.get('doc_count', 0)} 份 |",
        f"| 唯一学生 | {scan.get('student_count', 0)} 名 |",
        f"| 预警对象 | {display(warning_count, ' 位')} |",
        f"| 缺页漏页 | {display(missing_count, ' 位')} |",
        f"| 完成校对 | {display(completed_count, ' 位')} |",
        f"| 确认敏感学生 | {len({item['student'] for item in confirmed})} 名 |",
        f"| 确认敏感处数 | {sum(item.get('occurrences', 1) for item in confirmed)} 处 |",
        "",
    ]
    if counts_unknown:
        lines.extend([
            "> 预警/缺页人数未在 `review_meta.json` 中提供，报告中保留为“待确认”，未按其他班级数字套用。",
            "",
        ])
    if warning_names:
        lines.append(f"- 预警名单：{'、'.join(warning_names)}")
    if missing_names:
        lines.append(f"- 缺页漏页名单：{'、'.join(missing_names)}")
    if warning_names or missing_names:
        lines.append("")

    lines.extend([
        "## 二、敏感内容明细",
        "",
    ])
    if confirmed:
        lines.extend([
            "| 学生 | 类别 | 摘要 | 程度 | 处数 | 语境 |",
            "| --- | --- | --- | --- | --- | --- |",
        ])
        for item in confirmed:
            lines.append(
                f"| {item['student']} | {item['category_label']} | {item['summary']} | "
                f"{SEVERITY_LABELS.get(item['severity'], item['severity'])} | {item.get('occurrences', 1)} | "
                f"{item['context_type']} |"
            )
        lines.append("")
        lines.append("以上为结合上下文确认的敏感内容；未在汇报中展开原文。")
    else:
        lines.extend(["未确认需在汇报中单列的敏感内容。", ""])
    if excluded:
        lines.append(f"> 已排除关键词候选 {len(excluded)} 条（假设、转述、否定、比喻或一般性表述）；排除记录保留在本地 JSON。")
        lines.append("")

    lines.extend([
        "## 三、给老师的简短汇报",
        "",
        "```text",
        build_short_message(
            class_name, teacher, total, warning_count, missing_count, completed_count,
            confirmed, counts_unknown,
        ),
        "```",
        "",
        "## 四、口径与限制",
        "",
        f"- 统计口径来源：`{scan.get('meta', {}).get('count_source', 'unknown')}`",
        "- 本报告只读取审核后的 Word，未修改、移动或重新生成学生文件。",
        "- 预警/缺页漏页仅按老师提供的信息统计；未提供名单时报告不推断姓名。",
        "- 敏感内容由关键词候选加前后文判断，不替代老师的最终判断。",
        "- 敏感原文只保留在本地报告中，不在聊天或外部服务中展开。",
    ])
    if warnings:
        lines.extend(["", "**数据校验提示：**"])
        lines.extend(f"- {item}" for item in warnings)
    return "\n".join(lines) + "\n"


def build_docx(path: Path, class_name, teacher, total, warning_count, missing_count,
               completed_count, warning_names, missing_names, confirmed, excluded,
               counts_unknown, warnings, scan, decisions):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5

    add_text(doc, f"{class_name}自传审核汇总报告", east="黑体", size=16, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_text(doc, f"审核对象：{scan.get('input_dir', '')}", size=10.5,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

    add_heading(doc, "一、总体情况")
    rows = [
        ["审核文件", f"{scan.get('doc_count', 0)} 份"],
        ["唯一学生", f"{scan.get('student_count', 0)} 名"],
        ["预警对象", display(warning_count, " 位")],
        ["缺页漏页", display(missing_count, " 位")],
        ["完成校对", display(completed_count, " 位")],
        ["确认敏感学生", f"{len({item['student'] for item in confirmed})} 名"],
        ["确认敏感处数", f"{sum(item.get('occurrences', 1) for item in confirmed)} 处"],
    ]
    add_table(doc, ["项目", "数量"], rows, widths=[7.0, 8.0])
    if counts_unknown:
        add_text(doc, "预警/缺页人数未在 review_meta.json 中提供，报告中保留为“待确认”，未套用其他班级数字。",
                 size=10.5, before=5, after=8)
    if warning_names:
        add_text(doc, f"预警名单：{'、'.join(warning_names)}", size=10.5, after=3)
    if missing_names:
        add_text(doc, f"缺页漏页名单：{'、'.join(missing_names)}", size=10.5, after=8)

    add_heading(doc, "二、敏感内容明细")
    if confirmed:
        rows = [
            [
                item["student"],
                item["category_label"],
                item["summary"],
                SEVERITY_LABELS.get(item["severity"], item["severity"]),
                str(item.get("occurrences", 1)),
                item["context_type"],
            ]
            for item in confirmed
        ]
        add_table(doc, ["学生", "类别", "摘要", "程度", "处数", "语境"], rows,
                  widths=[2.2, 3.0, 4.6, 1.6, 1.2, 2.8])
        add_text(doc, "以上为结合上下文确认的敏感内容；报告不展开原文。", size=10.5, before=5, after=8)
    else:
        add_text(doc, "未确认需在汇报中单列的敏感内容。", after=8)
    if excluded:
        add_text(doc, f"已排除关键词候选 {len(excluded)} 条（假设、转述、否定、比喻或一般性表述）。",
                 size=10.5, after=8)

    add_heading(doc, "三、给老师的简短汇报")
    for line in build_short_message(
        class_name, teacher, total, warning_count, missing_count, completed_count,
        confirmed, counts_unknown,
    ).splitlines():
        if line.strip():
            add_text(doc, line, after=3)

    add_heading(doc, "四、口径与限制")
    for line in (
        f"统计口径来源：{scan.get('meta', {}).get('count_source', 'unknown')}",
        "本报告只读取审核后的 Word，未修改、移动或重新生成学生文件。",
        "预警/缺页漏页仅按老师提供的信息统计；未提供名单时不推断姓名。",
        "敏感内容由关键词候选加前后文判断，不替代老师的最终判断。",
        "敏感原文只保留在本地报告中，不在聊天或外部服务中展开。",
    ):
        add_text(doc, "• " + line, after=3)
    if warnings:
        add_text(doc, "数据校验提示：", east="黑体", bold=True, before=6, after=2)
        for item in warnings:
            add_text(doc, "• " + item, size=10.5, after=2)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a teacher-facing report from scanner and decisions.")
    parser.add_argument("--scan", required=True)
    parser.add_argument("--decisions", required=True)
    parser.add_argument("--out", required=True)
    parser.add_argument("--class-name", default="")
    parser.add_argument("--teacher", default="老师")
    parser.add_argument("--report-name", default="审核汇总报告")
    args = parser.parse_args()

    scan_path = Path(args.scan).expanduser().resolve()
    decisions_path = Path(args.decisions).expanduser().resolve()
    out_dir = Path(args.out).expanduser().resolve()
    scan = load_json(scan_path)
    decisions = load_json(decisions_path)
    confirmed, excluded, warnings = normalize_decisions(decisions, scan)
    meta = scan.get("meta") or {}
    class_name = args.class_name or meta.get("class_name") or "班级"
    total = meta.get("total_students")
    if total is None:
        total = scan.get("student_count")
    warning_count = meta.get("warning_count")
    missing_count = meta.get("missing_pages_count")
    completed_count = meta.get("completed_count")
    if completed_count is None and total is not None and missing_count is not None:
        completed_count = int(total) - int(missing_count)
    counts_unknown = any(value is None for value in (total, warning_count, missing_count, completed_count))
    warnings.extend(scan.get("meta_warnings") or [])

    if not decisions.get("reviewed"):
        warnings.append("sensitive_decisions.json 的 reviewed 不是 true：当前敏感结论未完成复核")
    warnings = list(dict.fromkeys(warnings))

    out_dir.mkdir(parents=True, exist_ok=True)
    md_path = out_dir / f"{args.report_name}.md"
    docx_path = out_dir / f"{args.report_name}.docx"
    txt_path = out_dir / "发给老师的汇报.txt"

    md = build_markdown(
        class_name, args.teacher, total, warning_count, missing_count, completed_count,
        meta.get("warning") or [], meta.get("missing_pages") or [], confirmed, excluded,
        counts_unknown, warnings, scan, decisions,
    )
    md_path.write_text(md, encoding="utf-8-sig", newline="\n")

    build_docx(
        docx_path, class_name, args.teacher, total, warning_count, missing_count,
        completed_count, meta.get("warning") or [], meta.get("missing_pages") or [],
        confirmed, excluded, counts_unknown, warnings, scan, decisions,
    )

    txt = build_short_message(
        class_name, args.teacher, total, warning_count, missing_count,
        completed_count, confirmed, counts_unknown,
    )
    txt_path.write_text(txt, encoding="utf-8-sig", newline="\n")

    print(f"confirmed={len(confirmed)} students={len({item['student'] for item in confirmed})} excluded={len(excluded)}")
    print(str(docx_path))
    print(str(md_path))
    print(str(txt_path))
    if warnings:
        print("warnings=" + " | ".join(warnings))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
